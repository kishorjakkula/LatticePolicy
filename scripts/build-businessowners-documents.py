#!/usr/bin/env python3
"""Build original BOP merge templates with the bundled python-docx runtime.

Run with the Python path returned by Codex load_workspace_dependencies.
No customer data, external requests or database access are used.
"""
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'products' / 'businessowners' / 'documents'
WARNING = 'DEMO - NOT A CONTRACT OF INSURANCE'


def paragraph(document, text, style=None):
    return document.add_paragraph(text, style)


def facts(document, rows):
    table = document.add_table(rows=0, cols=2)
    table.autofit = False
    table.columns[0].width = Inches(1.6)
    table.columns[1].width = Inches(5.4)
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].width = Inches(1.6)
        cells[1].width = Inches(5.4)
        cells[0].text = label
        cells[1].text = value
        for run in cells[0].paragraphs[0].runs:
            run.bold = True
        props = table.rows[-1]._tr.get_or_add_trPr()
        props.append(OxmlElement('w:cantSplit'))
    paragraph(document, '')


def build(kind, title, intro):
    document = Document()
    document.core_properties.author = 'LatticePolicy'
    document.core_properties.title = title
    document.core_properties.subject = 'Original synthetic BOP demonstration document'
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.65)
    section.left_margin = section.right_margin = Inches(.75)
    for name in ['Normal', 'Title', 'Heading 1', 'Heading 2']:
        style = document.styles[name]
        style.font.name = 'Arial'
        style.font.color.rgb = RGBColor.from_string('000000')
        style.paragraph_format.space_after = Pt(6)
        borders = style.element.xpath('./w:pPr/w:pBdr')
        for border in borders:
            border.getparent().remove(border)
    normal = document.styles['Normal']
    normal.font.size = Pt(10)
    normal.paragraph_format.line_spacing = 1.05
    document.styles['Title'].font.size = Pt(22)
    document.styles['Heading 1'].font.size = Pt(13)
    footer = section.footer.paragraphs[0]
    footer.text = f'{WARNING}  |  Original LatticePolicy demo  |  {kind}'
    footer.runs[0].font.size = Pt(8)
    paragraph(document, title, 'Title')
    warning = paragraph(document, WARNING)
    warning.runs[0].bold = True
    paragraph(document, intro)
    paragraph(document, 'Transaction record', 'Heading 1')
    facts(document, [
        ('Named insured', '{{insuredName}}'),
        ('Mailing address', '{{mailingAddress}}'),
        ('Demo insurer', '{{underwritingCompany}}'),
        ('Policy reference', '{{policyNumber}}'),
        ('Transaction', '{{transactionType}}  {{transactionNumber}}'),
        ('Period', '{{effectiveDate}} to {{expirationDate}}'),
        ('Jurisdiction', '{{state}}'),
        ('Agency', '{{agencyName}}'),
        ('Producer', '{{agentName}}'),
    ])
    paragraph(document, 'Business and premises', 'Heading 1')
    paragraph(document, '{{businessownersRiskSummary}}')
    paragraph(document, 'Demonstration premium', 'Heading 1')
    facts(document, [
        ('Premium', '{{premiumTotal}} {{premiumCurrency}}'),
        ('Fees', '{{premiumFees}} {{premiumCurrency}}'),
        ('Taxes', '{{premiumTaxes}} {{premiumCurrency}}'),
    ])
    paragraph(document, 'Amounts use synthetic non-filed demo pricing. A zero tax amount is not a tax exemption or a statement of applicable law.')
    document.add_page_break()
    paragraph(document, 'Selected coverages and package terms', 'Heading 1')
    paragraph(document, WARNING).runs[0].bold = True
    paragraph(document, 'The following summary reflects the selected transaction snapshot. Included benefits share the stated package limits where identified; they are not additional separately charged coverages.')
    paragraph(document, '{{businessownersCoverageSummary}}')
    paragraph(document, 'Scope and limitations', 'Heading 1')
    paragraph(document, 'This original demonstration is not an insurance offer, binder of actual coverage, policy wording, evidence of insurance, or approved filing. No ISO or other proprietary policy wording is reproduced. No representation of admission, regulatory approval or licensed distribution is made.')
    paragraph(document, 'No unrestricted actual loss sustained benefit is represented. Auto, workers compensation, professional liability, cyber, flood and earthquake are not automatically included in this demo package. Any real offering requires carrier and legal review of wording, rates, eligibility, exclusions, notices and distribution authority.')
    paragraph(document, 'Historical record', 'Heading 1')
    paragraph(document, 'Read this document with its recorded transaction and form edition. Later transactions or template changes do not amend this archived output. This summary does not grant, extend or certify coverage.')
    OUT.mkdir(parents=True, exist_ok=True)
    document.save(OUT / f'{kind}.docx')


if __name__ == '__main__':
    build('BOP-QUOTE', 'Businessowners Demo Quote',
          'Review the proposed business premises, selected package terms and demonstration premium below. This proposal creates no insurance coverage or obligation to bind.')
    build('BOP-BINDER', 'Businessowners Demo Binder',
          'This record demonstrates the bind transaction in the software only. It does not bind actual insurance, certify coverage in force or establish a real binder period.')
    build('BOP-DEC', 'Businessowners Demo Declarations',
          'This record summarizes the business premises, selected package terms and premium for the identified demonstration transaction. It is not a contract of insurance.')
